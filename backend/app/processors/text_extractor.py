import os
import json
import csv
import logging
from typing import Dict, Any
from pathlib import Path
from pypdf import PdfReader
import docx
from app.processors.ocr_engine import perform_ocr

logger = logging.getLogger("investigation.extractor")

def extract_text_from_file(file_path: str, original_filename: str) -> Dict[str, Any]:
    """
    Extracts text from diverse digital forensic evidence formats.
    Supported: PDF, DOCX, TXT, CSV, JSON, PNG, JPG, JPEG, TIFF, WEBP, LOG, EML.
    """
    ext = Path(original_filename).suffix.lower()
    result = {
        "extracted_text": "",
        "extraction_method": "unknown",
        "character_count": 0,
        "page_count": 1,
        "status": "success",
        "error": None
    }

    if not os.path.exists(file_path):
        result["status"] = "failed"
        result["error"] = "File not found on disk"
        return result

    try:
        # 1. Plain Text, Log files, EML, Markdown
        if ext in [".txt", ".log", ".chat", ".eml", ".md"]:
            with open(file_path, "r", encoding="utf-8", errors="replace") as f:
                content = f.read()
            result["extracted_text"] = content
            result["extraction_method"] = "plain_text"

        # 2. JSON files
        elif ext == ".json":
            with open(file_path, "r", encoding="utf-8", errors="replace") as f:
                data = json.load(f)
            result["extracted_text"] = json.dumps(data, indent=2)
            result["extraction_method"] = "json_parser"

        # 3. CSV files
        elif ext == ".csv":
            lines = []
            with open(file_path, "r", encoding="utf-8", errors="replace") as f:
                reader = csv.reader(f)
                for row in reader:
                    lines.append(" | ".join(row))
            result["extracted_text"] = "\n".join(lines)
            result["extraction_method"] = "csv_parser"

        # 4. DOCX Documents
        elif ext == ".docx":
            doc = docx.Document(file_path)
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_text:
                        paragraphs.append(" | ".join(row_text))
            result["extracted_text"] = "\n\n".join(paragraphs)
            result["extraction_method"] = "docx_parser"

        # 5. PDF Documents
        elif ext == ".pdf":
            reader = PdfReader(file_path)
            pages_text = []
            result["page_count"] = len(reader.pages)
            for i, page in enumerate(reader.pages):
                txt = page.extract_text()
                if txt and txt.strip():
                    pages_text.append(txt.strip())

            extracted = "\n\n".join(pages_text)
            if extracted.strip():
                result["extracted_text"] = extracted
                result["extraction_method"] = "pdf_native"
            else:
                # Scanned PDF fallback
                result["extraction_method"] = "pdf_scanned_ocr_needed"
                result["extracted_text"] = "[Scanned PDF - No embedded text stream found]"
                result["status"] = "partial"

        # 6. Images (OCR)
        elif ext in [".png", ".jpg", ".jpeg", ".tiff", ".bmp", ".webp"]:
            ocr_text, success, msg = perform_ocr(file_path)
            result["extracted_text"] = ocr_text
            result["extraction_method"] = "tesseract_ocr"
            if not success and not ocr_text:
                result["status"] = "partial"
                result["error"] = msg

        else:
            # Fallback binary text extraction
            with open(file_path, "rb") as f:
                raw = f.read(50000)
            text_chars = [chr(b) for b in raw if 32 <= b <= 126 or b in [10, 13, 9]]
            extracted = "".join(text_chars).strip()
            if extracted:
                result["extracted_text"] = extracted[:5000]
                result["extraction_method"] = "binary_strings"
            else:
                result["status"] = "failed"
                result["error"] = f"Unsupported file extension: {ext}"

    except Exception as e:
        logger.error(f"Error extracting text from {file_path}: {e}")
        result["status"] = "failed"
        result["error"] = str(e)

    result["character_count"] = len(result["extracted_text"])
    return result
